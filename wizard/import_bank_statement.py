# -*- coding: utf-8 -*-
###############################################################################
#
#    Cybrosys Technologies Pvt. Ltd.
#
#    Copyright (C) 2025-TODAY Cybrosys Technologies(<https://www.cybrosys.com>)
#    Author: Akhil Ashok (odoo@cybrosys.com)
#
#    You can modify it under the terms of the GNU LESSER
#    GENERAL PUBLIC LICENSE (LGPL v3), Version 3.
#
#    This program is distributed in the hope that it will be useful,
#    but WITHOUT ANY WARRANTY; without even the implied warranty of
#    MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
#    GNU LESSER GENERAL PUBLIC LICENSE (LGPL v3) for more details.
#
#    You should have received a copy of the GNU LESSER GENERAL PUBLIC LICENSE
#    (LGPL v3) along with this program.
#    If not, see <http://www.gnu.org/licenses/>.
#
###############################################################################
import base64
import codecs
import csv
import openpyxl
import os
import io
from datetime import datetime
from io import BytesIO
from odoo import fields, models, _
from odoo.exceptions import ValidationError
from ofxparse import OfxParser
from qifparse.parser import QifParser


class UniversalImportBankStatement(models.TransientModel):
    """ A class to import files as bank statement """
    _name = "universal.universal.import.bank.statement"
    _description = "Import button"
    _rec_name = "file_name"

    attachment = fields.Binary(string="File", required=True,
                               help="Choose the file to import")
    file_name = fields.Char(string="File Name", help="Name of the file")
    journal_id = fields.Many2one('account.journal', string="Journal ID",
                                 help="Journal in which the file importing")

    def _parse_date(self, date_str):
        """ Helper to parse date from string """
        if not date_str:
            return fields.Date.today()

        if isinstance(date_str, datetime):
            return date_str.date()

        # Remove potential quotes and whitespace
        date_str = str(date_str).strip().strip('"').strip("'")
        if not date_str:
            return fields.Date.today()

        # Try common formats
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(date_str, fmt).date()
            except (ValueError, TypeError):
                continue

        # Fallback to Odoo fields.Date.from_string
        try:
            res = fields.Date.from_string(date_str)
            if res:
                return res
        except:
            pass

        return fields.Date.today()

    def _parse_float(self, val):
        """ Helper to parse float from string with currency symbols and commas """
        if not val:
            return 0.0
        if isinstance(val, (int, float)):
            return float(val)
        
        # Remove currency symbols, quotes, commas, spaces
        clean_val = str(val).strip().replace('"', '').replace("'", "").replace(',', '').replace(' ', '')
        for symbol in ('$', '€', '£', '¥', '₹'):
            clean_val = clean_val.replace(symbol, '')

        # Handle accounting negative format (100.0) -> -100.0
        if clean_val.startswith('(') and clean_val.endswith(')'):
            clean_val = '-' + clean_val[1:-1]

        try:
            return float(clean_val)
        except (ValueError, TypeError):
            return 0.0

    def action_statement_import(self):
        """Function to import csv, xlsx, ofx and qif file format"""
        split_tup = os.path.splitext(self.file_name)
        if split_tup[1] == '.csv' or split_tup[1] == '.xlsx' or split_tup[
            1] == '.ofx' or split_tup[1] == '.qif':
            if split_tup[1] == '.csv':
                try:
                    file_data = base64.b64decode(self.attachment)
                    file_string = file_data.decode('utf-8-sig')
                    f = io.StringIO(file_string)
                    reader = csv.DictReader(f)
                    fieldnames = reader.fieldnames or []
                    header_map = {f.strip().lower(): f for f in fieldnames}
                except Exception as e:
                    raise ValidationError(_("Error reading CSV file: %s") % str(e))

                statement_id = False
                for row in reader:
                    if not any(row.values()):
                        continue

                    def get_field(keys):
                        for k in keys:
                            if k.lower() in header_map:
                                return row.get(header_map[k.lower()])
                        return None

                    # Mappings for common Odoo and generic bank headers
                    name = get_field(['name', 'label', 'description', 'reference', 'account', 'line_ids/payment_ref', 'payment_ref'])
                    amount = get_field(['amount', 'value', 'price', 'total', 'line_ids/amount'])
                    partner_name = get_field(['partner', 'partner_id/name', 'contact', 'payee', 'customer', 'supplier', 'line_ids/partner_id', 'partner_id'])
                    date_str = get_field(['date', 'transaction date', 'time', 'line_ids/date'])
                    starting_balance = get_field(['starting balance', 'start balance', 'opening balance', 'balance_start'])
                    ending_balance = get_field(['ending balance', 'balance', 'real balance', 'end balance', 'balance_end_real'])

                    values = [v.strip() if v else '' for v in row.values()]
                    keys = list(row.keys())
                    mapped_indices = set()
                    
                    def find_column(targets, is_numeric=False, is_date=False, exclude_indices=None):
                        exclude_indices = exclude_indices or set()
                        for k, v in header_map.items():
                            if any(t in k for t in targets):
                                idx = keys.index(v)
                                if idx not in exclude_indices:
                                    return row.get(v), idx
                        if is_numeric or is_date:
                            for i, val in enumerate(values):
                                if i in exclude_indices: continue
                                if is_numeric:
                                    try:
                                        temp = val.replace('$', '').replace(',', '').strip()
                                        if temp:
                                            float(temp)
                                            return val, i
                                    except: pass
                                if is_date:
                                    try:
                                        self._parse_date(val)
                                        return val, i
                                    except: pass
                        return None, -1

                    # Heuristic backups for unmapped fields
                    if amount is None:
                        amount_val, idx = find_column(['amount', 'value', 'price', 'total'], is_numeric=True)
                        if idx != -1:
                            amount = amount_val
                            mapped_indices.add(idx)

                    if date_str is None:
                        date_val, idx = find_column(['date', 'time'], is_date=True, exclude_indices=mapped_indices)
                        if idx != -1:
                            date_str = date_val
                            mapped_indices.add(idx)

                    if ending_balance is None:
                        bal_val, idx = find_column(['balance', 'ending', 'real', 'end balance'], is_numeric=True, exclude_indices=mapped_indices)
                        if idx != -1:
                            ending_balance = bal_val
                            mapped_indices.add(idx)

                    if starting_balance is None:
                        # Attempt to find starting balance only if multiple numeric columns exist
                        numeric_count = len([v for v in values if v.replace('$', '').replace(',', '').strip().replace('.', '').isdigit()])
                        bal_val, idx = find_column(['starting', 'start balance', 'opening'], is_numeric=(numeric_count > 2), exclude_indices=mapped_indices)
                        if idx != -1:
                            starting_balance = bal_val
                            mapped_indices.add(idx)

                    if name is None:
                        for i, val in enumerate(values):
                            if i in mapped_indices: continue
                            try:
                                float(val.replace('$', '').replace(',', '').strip())
                                continue
                            except:
                                try:
                                    self._parse_date(val)
                                    continue
                                except:
                                    name = val
                                    mapped_indices.add(i)
                                    break
                    
                    if not name and values:
                        name = values[0]

                    transaction_date = self._parse_date(date_str)
                    clean_amount = self._parse_float(amount)
                    clean_start_balance = self._parse_float(starting_balance)
                    clean_end_balance = self._parse_float(ending_balance)

                    # Ensure statement is balanced (End = Start + Amount) to prevent red state
                    if starting_balance is not None and ending_balance is not None:
                        b_start, b_end = clean_start_balance, clean_end_balance
                        clean_amount = b_end - b_start
                    elif ending_balance is not None:
                        b_end = clean_end_balance
                        b_start = b_end - clean_amount
                    elif starting_balance is not None:
                        b_start = clean_start_balance
                        b_end = b_start + clean_amount
                    else:
                        b_start, b_end = 0.0, clean_amount

                    partner = False
                    if partner_name:
                        partner_name = partner_name.strip()
                        if partner_name.lower() in ('bank', 'cash', 'main', 'demo', 'yourcompany'):
                            partner_name = False
                        if partner_name:
                            partner = self.env['res.partner'].search([('name', '=', partner_name)], limit=1)
                            if not partner and len(partner_name) > 1 and not partner_name.isdigit():
                                try:
                                    self._parse_date(partner_name)
                                except:
                                    raise ValidationError(_("Partner '%s' does not exist") % partner_name)

                    statement = self.env['account.bank.statement'].create({
                        'name': name,
                        'journal_id': self.journal_id.id,
                        'company_id': self.journal_id.company_id.id,
                        'date': transaction_date,
                        'balance_start': b_start,
                        'balance_end_real': b_end,
                        'line_ids': [(0, 0, {
                            'date': transaction_date,
                            'payment_ref': name or 'csv file',
                            'partner_id': partner.id if partner else False,
                            'journal_id': self.journal_id.id,
                            'amount': clean_amount,
                        })],
                    })
                    statement_id = statement.id

                return {
                    'type': 'ir.actions.act_window',
                    'name': 'Statements',
                    'view_mode': 'list',
                    'res_model': 'account.bank.statement',
                    'res_id': statement.id,
                }
            elif split_tup[1] == '.xlsx':
                # Reading xlsx file — supports flexible column layouts including
                # CBZ/CABS ZWG format: Posting Date | Description | Debits | Credits | Balance
                try:
                    order = openpyxl.load_workbook(
                        filename=BytesIO(base64.b64decode(self.attachment)),
                        data_only=True, read_only=True)
                    xl_order = order.active
                except Exception as e:
                    raise ValidationError(_("Choose correct file: %s") % str(e))

                # Detect header row and column positions (scan first 5 rows)
                headers = []
                header_row_idx = 1
                for i, row in enumerate(xl_order.iter_rows(max_row=5, values_only=True)):
                    row_vals = [str(c).strip().lower() if c else '' for c in row]
                    str_count = sum(1 for v in row_vals if v and not self._parse_float(v))
                    if str_count >= 2:
                        headers = row_vals
                        header_row_idx = i + 1
                        break

                def _find_col(keys, hdrs):
                    for key in keys:
                        for j, h in enumerate(hdrs):
                            if h and key in h:
                                return j
                    return -1

                col_date    = _find_col(['date', 'posting', 'tran date', 'value date'], headers)
                col_desc    = _find_col(['description', 'narration', 'particulars', 'reference', 'detail'], headers)
                col_debit   = _find_col(['debit', 'withdrawal', 'dr'], headers)
                col_credit  = _find_col(['credit', 'deposit', 'cr'], headers)
                col_amount  = _find_col(['amount', 'value'], headers)
                col_balance = _find_col(['balance', 'running balance'], headers)

                statement_lines = []
                start_balance = None
                end_balance = 0.0

                for record in xl_order.iter_rows(min_row=header_row_idx + 1,
                                                 values_only=True):
                    line = list(record)
                    if not any(line):
                        continue

                    def _get(col):
                        if col >= 0 and col < len(line):
                            return line[col]
                        return None

                    # --- Date ---
                    raw_date = _get(col_date) if col_date >= 0 else (line[0] if line else None)
                    from datetime import date as date_type
                    if isinstance(raw_date, (date_type, datetime)):
                        date_obj = raw_date.date() if isinstance(raw_date, datetime) else raw_date
                    elif isinstance(raw_date, (int, float)) and 40000 < raw_date < 60000:
                        from datetime import timedelta
                        excel_epoch = datetime(1899, 12, 30)
                        date_obj = (excel_epoch + timedelta(days=int(raw_date))).date()
                    else:
                        date_obj = self._parse_date(raw_date)

                    # --- Description / Label ---
                    desc = _get(col_desc)
                    if not desc:
                        for cell in line:
                            if cell and not isinstance(cell, (int, float, date_type, datetime)):
                                try:
                                    float(str(cell).replace(',', ''))
                                except (ValueError, TypeError):
                                    desc = str(cell)
                                    break
                    label = str(desc).strip() if desc else 'bank transaction'

                    # --- Amount: prefer debit/credit split columns ---
                    amount = 0.0
                    if col_debit >= 0 or col_credit >= 0:
                        debit  = self._parse_float(_get(col_debit))  if col_debit  >= 0 else 0.0
                        credit = self._parse_float(_get(col_credit)) if col_credit >= 0 else 0.0
                        amount = credit - debit
                    elif col_amount >= 0:
                        amount = self._parse_float(_get(col_amount))

                    if amount == 0.0:
                        continue

                    # --- Balance ---
                    balance_raw = _get(col_balance) if col_balance >= 0 else None
                    if balance_raw is not None:
                        current_bal = self._parse_float(balance_raw)
                        if start_balance is None:
                            start_balance = current_bal - amount
                        end_balance = current_bal

                    statement_lines.append((0, 0, {
                        'date': date_obj,
                        'payment_ref': label,
                        'journal_id': self.journal_id.id,
                        'amount': amount,
                    }))

                if not statement_lines:
                    raise ValidationError(_(
                        "No valid transactions found in the XLSX file. "
                        "Please ensure the file has columns for Date, Description, "
                        "and Debit/Credit or Amount."))

                if start_balance is None:
                    start_balance = 0.0
                    end_balance = sum(ln[2]['amount'] for ln in statement_lines)

                statement = self.env['account.bank.statement'].create({
                    'name': self.file_name or 'XLSX Import',
                    'journal_id': self.journal_id.id,
                    'company_id': self.journal_id.company_id.id,
                    'date': statement_lines[0][2]['date'] if statement_lines else fields.Date.context_today(self),
                    'balance_start': start_balance,
                    'balance_end_real': end_balance,
                    'line_ids': statement_lines,
                })

                return {
                    'type': 'ir.actions.client',
                    'tag': 'reload',
                }
            elif split_tup[1] == '.ofx':
                try:
                    file_data = base64.b64decode(self.attachment)
                    ofx_file = OfxParser.parse(io.BytesIO(file_data))
                except Exception as e:
                    raise ValidationError(_("Wrong file format or parsing error: %s") % str(e))

                if not ofx_file.account or not ofx_file.account.statement:
                    raise ValidationError(_("No account information found in OFX file."))

                statement_id = False
                stmt = ofx_file.account.statement
                
                # Standardize balance extraction
                ledger_bal = getattr(stmt, 'balance', getattr(stmt, 'ledger_balance', None))
                final_balance = self._parse_float(ledger_bal) if ledger_bal is not None else 0.0
                
                for transaction in stmt.transactions:
                    amount = self._parse_float(transaction.amount)
                    if amount == 0:
                        continue

                    # Clean labels and match partners
                    label = (transaction.memo or transaction.name or transaction.payee or 'ofx transaction').strip()
                    payee_name = transaction.payee.strip() if transaction.payee else False
                    
                    partner = False
                    if payee_name:
                        # Shared noise-filtering for partners
                        if payee_name.lower() in ('bank', 'cash', 'main', 'demo', 'yourcompany'):
                            payee_name = False
                        if payee_name:
                            partner = self.env['res.partner'].search([('name', '=', payee_name)], limit=1)
                    
                    date = self._parse_date(transaction.date)
                    b_end = final_balance
                    b_start = b_end - amount

                    statement = self.env['account.bank.statement'].create({
                        'name': label,
                        'journal_id': self.journal_id.id,
                        'company_id': self.journal_id.company_id.id,
                        'date': date,
                        'balance_start': b_start,
                        'balance_end_real': b_end,
                        'line_ids': [(0, 0, {
                            'date': date,
                            'payment_ref': label,
                            'partner_id': partner.id if partner else False,
                            'journal_id': self.journal_id.id,
                            'amount': amount,
                        })],
                    })
                    statement_id = statement.id
                
                if not statement_id:
                    raise ValidationError(_("No valid transactions found in the OFX file."))
                
                return {
                    'type': 'ir.actions.act_window',
                    'name': 'Statements',
                    'view_mode': 'list',
                    'res_model': 'account.bank.statement',
                    'res_id': statement_id,
                }

            elif split_tup[1] == '.qif':
                try:
                    file_data = base64.b64decode(self.attachment)
                    file_string = file_data.decode('utf-8-sig')
                    qif = QifParser().parse(io.StringIO(file_string))
                except Exception as e:
                    raise ValidationError(_("Error parsing QIF file: %s") % str(e))

                statement_id = False
                for account in qif.get_accounts():
                    for transaction in qif.get_transactions(account):
                        amount = self._parse_float(transaction.amount)
                        if amount == 0:
                            continue
                            
                        date = self._parse_date(transaction.date)
                        label = (transaction.payee or transaction.memo or 'qif transaction').strip()
                        payee_name = transaction.payee.strip() if transaction.payee else False
                        
                        partner = False
                        if payee_name:
                            if payee_name.lower() in ('bank', 'cash', 'main', 'demo', 'yourcompany'):
                                payee_name = False
                            if payee_name:
                                partner = self.env['res.partner'].search([('name', '=', payee_name)], limit=1)
                        
                        statement = self.env['account.bank.statement'].create({
                            'name': label,
                            'journal_id': self.journal_id.id,
                            'company_id': self.journal_id.company_id.id,
                            'date': date,
                            'balance_start': 0.0,
                            'balance_end_real': amount,
                            'line_ids': [(0, 0, {
                                'date': date,
                                'payment_ref': label,
                                'partner_id': partner.id if partner else False,
                                'journal_id': self.journal_id.id,
                                'amount': amount,
                            })],
                        })
                        statement_id = statement.id

                if not statement_id:
                    raise ValidationError(_("No valid transactions found in the QIF file."))

                return {
                    'type': 'ir.actions.act_window',
                    'name': 'Statements',
                    'view_mode': 'list',
                    'res_model': 'account.bank.statement',
                    'res_id': statement_id,
                }
        else:
            raise ValidationError(_("Choose correct file"))
